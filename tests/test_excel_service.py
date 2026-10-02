from pathlib import Path
import openpyxl
from bot.services.excel_service import generate_daily_excel, generate_monthly_excel


def test_generate_daily_excel(tmp_path: Path):
    summary_data = {
        "date": "2026-10-02",
        "total_hospitals": 16,
        "submitted_hospitals": 2,
        "total_breakfast": 240,
        "total_lunch": 250,
        "total_dinner": 230,
        "total_all": 720,
        "items": [
            {
                "region_name": "Qaraózek rayonı",
                "hospital_name": "Qaraózek RMO emlewxanası",
                "breakfast_count": 120,
                "lunch_count": 125,
                "dinner_count": 115,
                "total_daily": 360,
                "user_name": "Berdiyarov Azamat",
                "user_phone": "+998901234567",
                "updated_at": "2026-10-02 08:30:00",
                "note": "Rejim boyınsha"
            },
            {
                "region_name": "Taxtakópir rayonı",
                "hospital_name": "Taxtakópir RMO emlewxanası",
                "breakfast_count": 120,
                "lunch_count": 125,
                "dinner_count": 115,
                "total_daily": 360,
                "user_name": "Amanbaev Murat",
                "user_phone": "+998912345678",
                "updated_at": "2026-10-02 08:45:00",
                "note": ""
            }
        ]
    }

    out_file = tmp_path / "reports" / "daily_test.xlsx"
    res_path = generate_daily_excel(summary_data, out_file)
    assert res_path.exists()
    assert res_path == out_file

    wb = openpyxl.load_workbook(res_path)
    assert "Kúnlik esabat" in wb.sheetnames
    ws = wb["Kúnlik esabat"]

    # Check title header
    assert "LAKOMA" in str(ws["A1"].value).upper()
    # Check date subtitle
    assert "2026-10-02" in str(ws["A2"].value)

    # Check data rows and totals
    found_first_hospital = False
    found_total = False
    total_row_values = []

    for row in ws.iter_rows(values_only=True):
        if row and any(x and "Qaraózek RMO" in str(x) for x in row):
            found_first_hospital = True
            assert 120 in row
            assert 125 in row
            assert 115 in row
            assert 360 in row
        if row and any(x and "JÁMI" in str(x).upper() for x in row):
            found_total = True
            total_row_values = list(row)

    assert found_first_hospital is True
    assert found_total is True
    assert 240 in total_row_values
    assert 250 in total_row_values
    assert 230 in total_row_values
    assert 720 in total_row_values

    # Check column formatting/widths
    assert ws.column_dimensions["A"].width > 0
    assert ws.column_dimensions["C"].width >= 20


def test_generate_monthly_excel(tmp_path: Path):
    records = [
        {
            "report_date": "2026-10-01",
            "region_name": "Nókis qalası",
            "hospital_name": "Respublika kóp tarmaqlı emlewxanası",
            "breakfast_count": 100,
            "lunch_count": 110,
            "dinner_count": 95,
            "total_daily": 305,
            "user_name": "Xalmuratov Timur",
            "user_phone": "+998909876543",
            "note": "Oddiy"
        },
        {
            "report_date": "2026-10-02",
            "region_name": "Nókis qalası",
            "hospital_name": "Respublika kóp tarmaqlı emlewxanası",
            "breakfast_count": 105,
            "lunch_count": 115,
            "dinner_count": 100,
            "total_daily": 320,
            "user_name": "Xalmuratov Timur",
            "user_phone": "+998909876543",
            "note": ""
        }
    ]

    out_file = tmp_path / "reports" / "monthly_2026_10.xlsx"
    res_path = generate_monthly_excel(records, "2026-10", out_file)
    assert res_path.exists()
    assert res_path == out_file

    wb = openpyxl.load_workbook(res_path)
    sheet_name = f"Aylıq esabat 2026-10"
    assert sheet_name in wb.sheetnames
    ws = wb[sheet_name]

    # Check title header
    assert "LAKOMA" in str(ws["A1"].value).upper()
    assert "2026-10" in str(ws["A1"].value)

    # Check data rows and totals
    found_date_row = False
    found_total = False
    total_row_values = []

    for row in ws.iter_rows(values_only=True):
        if row and any(x and "2026-10-01" in str(x) for x in row):
            found_date_row = True
            assert 100 in row
            assert 110 in row
            assert 95 in row
            assert 305 in row
        if row and any(x and "JÁMI" in str(x).upper() for x in row):
            found_total = True
            total_row_values = list(row)

    assert found_date_row is True
    assert found_total is True
    assert 205 in total_row_values  # 100 + 105
    assert 225 in total_row_values  # 110 + 115
    assert 195 in total_row_values  # 95 + 100
    assert 625 in total_row_values  # 305 + 320

    # Check column formatting/widths
    assert ws.column_dimensions["A"].width > 0
    assert ws.column_dimensions["C"].width >= 20
