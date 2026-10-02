from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_daily_excel(summary_data: dict, output_path: Path | str) -> Path:
    target_path = Path(output_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kúnlik esabat"

    # Brand Colors: Lakoma Green / Medium Green / Gold
    header_fill = PatternFill(start_color="1E5631", end_color="1E5631", fill_type="solid")
    subheader_fill = PatternFill(start_color="4C9A2A", end_color="4C9A2A", fill_type="solid")
    total_fill = PatternFill(start_color="F4D03F", end_color="F4D03F", fill_type="solid")

    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    info_font = Font(name="Calibri", size=11, bold=True, color="333333")
    col_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    bold_data_font = Font(name="Calibri", size=11, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Title row
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "«LAKOMA» KATERING XÍZMETI - EMLEWXANALAR AWQAT ESABATÍ"
    title_cell.font = title_font
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitle with Date
    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Sáne: {summary_data.get('date', '')} | Qabıllandı: {summary_data.get('submitted_hospitals', 0)} / {summary_data.get('total_hospitals', 0)} emlewxana"
    sub_cell.font = info_font
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    # Table Headers
    headers = [
        "№", "Rayon", "Emlewxana (Bolnica) atı",
        "Azanǵı awqat", "Túslik (obet)", "Keshki awqat",
        "Kúnlik Jámi", "Juwapker xızmetker", "Telefon", "Eskertiw"
    ]
    
    ws.append([])  # row 3 empty spacer
    ws.row_dimensions[3].height = 10
    ws.row_dimensions[4].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.font = col_header_font
        cell.fill = subheader_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows
    current_row = 5
    items = summary_data.get("items", [])
    for idx, item in enumerate(items, 1):
        ws.row_dimensions[current_row].height = 22
        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2, value=item.get("region_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=3, value=item.get("hospital_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        
        c_br = ws.cell(row=current_row, column=4, value=item.get("breakfast_count", 0))
        c_br.alignment = Alignment(horizontal="right", vertical="center")
        c_lu = ws.cell(row=current_row, column=5, value=item.get("lunch_count", 0))
        c_lu.alignment = Alignment(horizontal="right", vertical="center")
        c_di = ws.cell(row=current_row, column=6, value=item.get("dinner_count", 0))
        c_di.alignment = Alignment(horizontal="right", vertical="center")
        c_tot = ws.cell(row=current_row, column=7, value=item.get("total_daily", 0))
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        c_tot.font = bold_data_font

        ws.cell(row=current_row, column=8, value=item.get("user_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=9, value=item.get("user_phone", "")).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=10, value=item.get("note", "")).alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 11):
            cell = ws.cell(row=current_row, column=c)
            cell.border = thin_border
            if c != 7:
                cell.font = data_font

        current_row += 1

    # Totals Row
    ws.row_dimensions[current_row].height = 26
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    total_label = ws.cell(row=current_row, column=1, value="ULÍWMA JÁMI (PORTCIYALAR):")
    total_label.font = bold_data_font
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    
    ws.cell(row=current_row, column=4, value=summary_data.get("total_breakfast", 0)).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=5, value=summary_data.get("total_lunch", 0)).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=6, value=summary_data.get("total_dinner", 0)).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=7, value=summary_data.get("total_all", 0)).alignment = Alignment(horizontal="right", vertical="center")

    for c in range(1, 11):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = total_fill
        cell.font = bold_data_font
        cell.border = thin_border

    # Adjust Column Widths (ignoring merged rows 1 & 2 for width calculation)
    for col in ws.iter_cols(min_row=4, max_row=current_row):
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['H'].width = 24
    ws.column_dimensions['I'].width = 18
    ws.column_dimensions['J'].width = 20

    wb.save(target_path)
    return target_path


def generate_monthly_excel(records: list[dict], month_str: str, output_path: Path | str) -> Path:
    target_path = Path(output_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Aylıq esabat {month_str}"

    header_fill = PatternFill(start_color="1E5631", end_color="1E5631", fill_type="solid")
    subheader_fill = PatternFill(start_color="4C9A2A", end_color="4C9A2A", fill_type="solid")
    total_fill = PatternFill(start_color="F4D03F", end_color="F4D03F", fill_type="solid")

    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"«LAKOMA» KATERING - AYLIQ ESABAT ({month_str})"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    headers = [
        "Sáne", "Rayon", "Emlewxana atı",
        "Azan", "Obet", "Kesh", "Kúnlik jámi", "Xızmetker", "Telefon"
    ]
    ws.append([])  # row 2 spacer
    ws.row_dimensions[2].height = 10
    ws.row_dimensions[3].height = 26
    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num, value=h)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = subheader_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    current_row = 4
    total_b = total_l = total_d = total_all = 0
    for r in records:
        ws.row_dimensions[current_row].height = 20
        ws.cell(row=current_row, column=1, value=r.get("report_date", "")).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2, value=r.get("region_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=3, value=r.get("hospital_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        
        b = r.get("breakfast_count", 0)
        l = r.get("lunch_count", 0)
        d = r.get("dinner_count", 0)
        t = r.get("total_daily", 0)
        total_b += b
        total_l += l
        total_d += d
        total_all += t

        ws.cell(row=current_row, column=4, value=b).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=current_row, column=5, value=l).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=current_row, column=6, value=d).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=current_row, column=7, value=t).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=current_row, column=8, value=r.get("user_name", "")).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=9, value=r.get("user_phone", "")).alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 10):
            ws.cell(row=current_row, column=c).border = thin_border
            ws.cell(row=current_row, column=c).font = Font(name="Calibri", size=11)
        ws.cell(row=current_row, column=7).font = Font(name="Calibri", size=11, bold=True)
        current_row += 1

    # Totals
    ws.row_dimensions[current_row].height = 26
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    total_lbl = ws.cell(row=current_row, column=1, value="ULÍWMA JÁMI:")
    total_lbl.alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=4, value=total_b).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=5, value=total_l).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=6, value=total_d).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(row=current_row, column=7, value=total_all).alignment = Alignment(horizontal="right", vertical="center")

    for c in range(1, 10):
        cell = ws.cell(row=current_row, column=c)
        cell.fill = total_fill
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.border = thin_border

    for col in ws.iter_cols(min_row=3, max_row=current_row):
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 34
    ws.column_dimensions['H'].width = 24
    ws.column_dimensions['I'].width = 18

    wb.save(target_path)
    return target_path
